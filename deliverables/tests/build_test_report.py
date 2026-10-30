"""System Test Report workbook (SEP490 Report 5 layout) from the JSON of run_system_test.py.

usage: build_test_report.py TEMPLATE.xlsx ROUND1.json OUT.xlsx [ROUND2.json [ROUND3.json]]
TEMPLATE is the blank SEP490 "Report5_Test Report.xlsx": sheet "Feature 1" is the style source of
every feature sheet.
"""
import json
import sys
from copy import copy
from datetime import date

import openpyxl
from openpyxl.styles import Alignment

tpl, src, out = sys.argv[1:4]
later = [{x['id']: x for x in json.load(open(f))} for f in sys.argv[4:6]]  # rounds 2 and 3, keyed by case id
PROJECT = 'Design and Evaluation of an AI Data Agent for Automated Data Analysis Using User-Provided Datasets'
CODE, TESTER = 'FA26SE115', 'Ngô Thành Đạt'
ENV = ('1. Stack: docker compose (deploy/docker-compose.dev.yml), logos commit af201866: management API on '
       'localhost:8080, PostgreSQL 16, engines dataset, hypothesis, experiment-design, execution, paper, oversight.\n'
       '2. Email: local log sender; sign-in codes are read from the management log.\n'
       '3. Client: Python requests script (scripts/run_system_test.py) on macOS 15; Origin http://localhost:3000.\n'
       '4. Accounts: a new Project Manager, Researcher and outsider per round; default Platform Admin admin@example.com.\n'
       '5. Data: a 60-row CSV, the local PostgreSQL as an external source, a Slack webhook URL (not called).\n'
       '6. LLM providers: none configured, so runs cannot start (cases marked Pending).')
REQ = {
    'Authentication & Account': 'Email-code sign-in, session, profile and sign-out (FE-01; UC-01).',
    'Project Workspace': 'Projects, members, invitations, settings, activity, archive, ownership and delete (FE-02; UC-03, UC-04).',
    'Data Management': 'Dataset upload, versions, preview, download and access rules (FE-03; UC-06).',
    'Data Connections & Integrations': 'External data connections and webhook integrations (FE-03; UC-07).',
    'Local Node': 'Pairing and revoking a Local Node runner (UC-08).',
    'Research Run Orchestration': 'Start, monitor, answer, accept, amend, stop and export a research run (FE-04; UC-09 to UC-13).',
    'Notifications': 'Notification preferences and in-app notifications.',
    'Administration': 'Platform Admin overview, usage and settings (UC-02, UC-05).',
}

rows = json.load(open(src))
features = list(dict.fromkeys(r['feature'] for r in rows))
when = date.fromisoformat(rows[0]['date'])
wb = openpyxl.load_workbook(tpl)
proto = wb['Feature 1']
S = {k: copy(proto[c]._style) for k, c in dict(func='A11', id='A12', text='A12', res='F12', plain='G12').items()}


from openpyxl.styles import PatternFill
FUNC_FILL = PatternFill('solid', fgColor='FFCCFFFF')


def put(ws, r, c, v, style):
    cell = ws.cell(row=r, column=c, value=v)
    cell._style = copy(S[style])
    return cell


# --- Cover
cv = wb['Cover']
cv['B4'], cv['B5'], cv['F4'], cv['F5'], cv['F6'] = PROJECT, CODE, TESTER, when, 'v1.0'
cv['B6'] = f'{CODE}_SystemTestReport_v1.0'
cv['A11'], cv['B11'], cv['C11'], cv['D11'] = when, 'v1.0', 'Initial system test report', 'A'
cv['E11'] = f'Rounds 1-{1 + len(later)} of {len(rows)} system test cases on the local stack'
cv['F11'] = 'Report III (SRS), Report V section 4.2, logos commit af201866'

# --- Test Cases list (one row per function)
tc = wb['Test Cases']
tc['D3'], tc['D4'], tc['D5'] = '=Cover!B4', '=Cover!B5', ENV
st_row = [copy(tc.cell(row=9, column=c)._style) for c in range(2, 7)]
for row in tc.iter_rows(min_row=9, max_row=tc.max_row):
    for c in row:
        c.value, c.hyperlink = None, None
funcs = list(dict.fromkeys((r['feature'], r['function']) for r in rows))
for i, (feat, fn) in enumerate(funcs, 1):
    pre = next(r['pre'] for r in rows if r['function'] == fn)
    for col, v in enumerate([i, fn, feat[:31], f'Test cases of "{fn}"', pre], 2):
        tc.cell(row=8 + i, column=col, value=v)._style = copy(st_row[col - 2])
    tc.cell(row=8 + i, column=4).hyperlink = f"#'{feat[:31]}'!A1"

# --- one sheet per feature
for feat in features:
    ws = wb.copy_worksheet(proto)
    ws.title = feat[:31]
    for row in ws.iter_rows(min_row=11, max_row=ws.max_row):
        for c in row:
            c.value = None
            c._style = copy(S['plain']) if c.column <= 15 else c._style
    for k in range(11, ws.max_row + 1):
        ws.row_dimensions[k].height = None
    ws['B2'], ws['B3'] = feat, REQ[feat]
    r = 11
    mine = [x for x in rows if x['feature'] == feat]
    for fn in dict.fromkeys(x['function'] for x in mine):
        for col in range(1, 16):  # function band spans the whole table, as in the template
            put(ws, r, col, fn if col == 1 else None, 'func').fill = copy(FUNC_FILL)
        r += 1
        for x in (y for y in mine if y['function'] == fn):
            put(ws, r, 1, x['id'], 'id')
            for col, v in ((2, x['desc']), (3, x['procedure']), (4, x['expected']), (5, x['pre'])):
                put(ws, r, col, v, 'text').alignment = Alignment(wrap_text=True, vertical='top')
            for col in range(9 + 3 * len(later), 15):
                put(ws, r, col, None, 'plain')
            put(ws, r, 6, x['result'], 'res')
            put(ws, r, 7, when, 'plain').number_format = 'mm/dd/yyyy'
            put(ws, r, 8, TESTER, 'plain')
            note = x['note']
            for k, rnd in enumerate(later, 1):  # round k+1 fills its result, date and tester columns
                y = rnd[x['id']]
                put(ws, r, 6 + 3 * k, y['result'], 'res')
                put(ws, r, 7 + 3 * k, date.fromisoformat(y['date']), 'plain').number_format = 'mm/dd/yyyy'
                put(ws, r, 8 + 3 * k, TESTER, 'plain')
                note = y['note']
            put(ws, r, 15, note, 'text').alignment = Alignment(wrap_text=True, vertical='top')
            r += 1
    last = r - 1
    ws['B4'] = f'=COUNTIF(A12:A{last},"ST-*")'
    for k, col in ((6, 'F'), (7, 'I'), (8, 'L')):  # each round counts its own column
        for j, L in enumerate('BCDE'):
            ws[f'{L}{k}'] = f'=COUNTIF(${col}$12:${col}${last},{L}$5)'
    ws.column_dimensions['O'].width = 48

# --- Test Statistics
sx = wb['Test Statistics']
sx['C3'], sx['C4'], sx['G3'], sx['H5'] = PROJECT, CODE, TESTER, when
sx['C5'] = f'{CODE}_SystemTestReport_v1.0'
sx['C6'] = f'Figures of round {1 + len(later)} on the local stack. Research Run cases that need an LLM provider are Pending.'
style = [copy(sx.cell(row=11, column=c)._style) for c in range(2, 9)]
for row in sx.iter_rows(min_row=11, max_row=sx.max_row):
    for c in row:
        c.value, c.hyperlink = None, None
for i, feat in enumerate(features):
    s = f"'{feat[:31]}'"
    k = 6 + len(later)  # statistics show the latest round
    vals = [i + 1, f'={s}!B2', f'={s}!B{k}', f'={s}!C{k}', f'={s}!D{k}', f'={s}!E{k}', f'={s}!B4']
    for col, v in enumerate(vals, 2):
        sx.cell(row=11 + i, column=col, value=v)._style = copy(style[col - 2])
end = 10 + len(features)
t = end + 2
sx.cell(row=t, column=3, value='Sub total')
for col in 'DEFGH':
    sx[f'{col}{t}'] = f'=SUM({col}11:{col}{end})'
sx[f'C{t + 2}'], sx[f'E{t + 2}'], sx[f'F{t + 2}'] = 'Test coverage', f'=(D{t}+E{t})*100/(H{t}-G{t})', '%'
sx[f'C{t + 3}'], sx[f'E{t + 3}'], sx[f'F{t + 3}'] = 'Test successful coverage', f'=D{t}*100/(H{t}-G{t})', '%'
for k in (t + 2, t + 3):
    sx[f'E{k}'].number_format = '0.0'

del wb['Feature 1'], wb['Feature 2']
wb.save(out)
print(len(features), 'features', len(rows), 'cases', {k: sum(r['result'] == k for r in rows) for k in ('Passed', 'Failed', 'Pending')})
