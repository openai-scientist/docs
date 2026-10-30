"""Unit Test workbook (SEP490 Report 5 layout) from unit_spec.py and a pytest JUnit XML.

usage: build_unit_test.py TEMPLATE.xlsx JUNIT.xml OUT.xlsx
TEMPLATE is a filled SEP490 unit test workbook: its Guideline, Cover, Functions, Statistics and one
function sheet give the styles; their sample content is cleared.
"""
import re
import sys
import xml.etree.ElementTree as ET
from copy import copy
from datetime import datetime

import openpyxl
from openpyxl.utils import get_column_letter

from unit_spec import FUNCTIONS
from unit_spec_more import MORE

FUNCTIONS = FUNCTIONS + MORE

tpl, junit, out = sys.argv[1:4]
PROJECT = 'Design and Evaluation of an AI Data Agent for Automated Data Analysis Using User-Provided Datasets'
CODE, CREATOR = 'FA26SE115', 'Ngô Thành Đạt'
ENV = ('1. Operating System: macOS 15 (Apple Silicon).\n'
       '2. Backend: Python 3.12, FastAPI, uv workspace "logos"; tests run with pytest and pytest-xdist (4 workers).\n'
       '3. Database: SQLite in-memory app per test; PostgreSQL tests (marker db) skip without PLATFORM_POSTGRES_TEST_URL.\n'
       '4. Command: uv run --package management pytest -n 4 -m "not live and not bench" --junitxml=management.xml\n'
       '5. External services (LLM engines, email, R2 storage, Google) are replaced by test doubles.\n'
       '6. Source: logos/services/management, commit af201866.')

# results of the real run, by (classname suffix, test function); a parametrised test passes only if every case passed
res, when = {}, None
root = ET.parse(junit).getroot()
for s in root.iter('testsuite'):
    when = when or datetime.fromisoformat(s.get('timestamp')).replace(tzinfo=None)
for tc in root.iter('testcase'):
    key = (tc.get('classname'), re.sub(r'\[.*', '', tc.get('name')))
    o = 'F' if tc.find('failure') is not None or tc.find('error') is not None else 'U' if tc.find('skipped') is not None else 'P'
    res[key] = o if res.get(key, 'P') == 'P' else res[key]


def result(file, test):
    hits = [o for (c, t), o in res.items() if c.endswith(file) and t == test]
    if not hits:
        sys.exit(f'no result for {file}::{test}')
    return 'F' if 'F' in hits else 'U' if 'U' in hits else 'P'


wb = openpyxl.load_workbook(tpl)
keep = {'Guideline', 'Cover', 'Functions', 'Statistics'}
proto = wb['importStock']
for name in wb.sheetnames:
    if name not in keep and wb[name] is not proto:
        del wb[name]

# styles of the template function sheet, taken before it is cleared
ST = {k: proto[c]._style for k, c in dict(head='A10', side='A11', nav='B9', pre='B10', grp='B21', val='D12', strip='E12',
                                          mark='F11', utc='F9', res='F43', label='B44', edge='E44', date='F45',
                                          blank='Z200').items()}


def put(ws, row, col, value, style):
    c = ws.cell(row=row, column=col, value=value)
    c._style = copy(ST[style])
    return c


def clear_from(ws, first):
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= first:
            ws.unmerge_cells(str(rng))
    for row in ws.iter_rows(min_row=first, max_row=ws.max_row):
        for c in row:
            c.value = None
            c._style = copy(ST['blank'])


# --- Cover
cv = wb['Cover']
cv['B4'], cv['F4'], cv['B5'], cv['F5'] = PROJECT, CREATOR, CODE, when
cv['B6'], cv['F6'] = f'{CODE}_UnitTest_v1.0', 'v1.0'
for row in cv.iter_rows(min_row=11, max_row=cv.max_row):
    for c in row:
        c.value = None
cv['A11'], cv['B11'], cv['C11'], cv['D11'] = when, 'v1.0', 'Initial unit test document', 'A'
cv['E11'] = f'Unit test cases for {len(FUNCTIONS)} management functions, results from the pytest run of {when:%Y-%m-%d}'
cv['F11'] = 'logos/services/management (tests/unit, tests/integration)'

# --- Functions list
fn = wb['Functions']
fn['E6'], fn['E7'] = 100, ENV
for row in fn.iter_rows(min_row=11, max_row=fn.max_row):
    for c in row:
        c.value, c.hyperlink = None, None
for i, f in enumerate(FUNCTIONS, 1):
    r = 10 + i
    vals = [i, f['req'], f['cls'], f['code'], f['code'], f['code'], f['desc'], f['pre']]
    for col, v in enumerate(vals, 1):
        fn.cell(row=r, column=col, value=v)._style = copy(fn.cell(row=11, column=col)._style)
    fn.cell(row=r, column=6).hyperlink = f"#'{f['code']}'!A1"

# --- one sheet per function
for i, f in enumerate(FUNCTIONS):
    ws = wb.copy_worksheet(proto)
    ws.title = f['code'][:31]
    clear_from(ws, 9)
    ws['C2'] = f"=Functions!E{11 + i}"
    ws['L2'], ws['C3'], ws['L3'], ws['C4'] = f['code'], CREATOR, CREATOR, f['loc']
    ws['L4'] = f'=MAX(0,ROUNDUP(C4*Functions!$E$6/1000,0)-O7)'
    ws['C5'] = f"{f['desc']} Test file: tests/…/{f['file'].replace('.', '/')}.py"
    cases = f['cases']
    n = len(cases)
    first, last = get_column_letter(6), get_column_letter(5 + n)
    for j in range(n):
        put(ws, 9, 6 + j, f'UTCID{j + 1:02d}', 'utc')
    for col in range(1, 6):
        put(ws, 9, col, None, 'nav')

    def value_row(text):
        # value text spans B:D so a long condition wraps instead of running over the group labels
        put(ws, row, 2, text, 'val').alignment = openpyxl.styles.Alignment(horizontal='right', vertical='center', wrap_text=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.row_dimensions[row].height = 12 * max(1, -(-len(text) // 34))

    row = 10
    put(ws, row, 2, 'Precondition', 'pre')
    row += 1
    value_row(f['pre'])
    for j in range(n):
        put(ws, row, 6 + j, 'O', 'mark')
    row += 1

    def block(kind):
        global row
        groups = {}
        for j, case in enumerate(cases):
            for item in case[kind]:
                g, v = item.split('|', 1)
                groups.setdefault(g, {}).setdefault(v, set()).add(j)
        for g, values in groups.items():
            put(ws, row, 2, g, 'grp')
            row += 1
            for v, cols in values.items():
                value_row(v)
                for j in range(n):
                    put(ws, row, 6 + j, 'O' if j in cols else None, 'mark')
                row += 1

    block(2)
    confirm_row = row
    block(3)
    row += 1
    for r in range(10, row):
        put(ws, r, 1, None, 'side'); put(ws, r, 5, None, 'strip')
    put(ws, 10, 1, 'Condition', 'head'); put(ws, confirm_row, 1, 'Confirm', 'head')
    for r in range(row, row + 5):
        put(ws, r, 1, None, 'side'); put(ws, r, 5, None, 'edge')
    put(ws, row, 1, 'Result', 'head'); put(ws, row, 2, 'Type(N : Normal, A : Abnormal, B : Boundary)', 'label')
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    trow = row
    for j, case in enumerate(cases):
        put(ws, row, 6 + j, case[1], 'res')
    row += 1
    put(ws, row, 2, 'Passed/Failed', 'label')
    prow = row
    for j, case in enumerate(cases):
        put(ws, row, 6 + j, result(f['file'], case[0]), 'res')
    row += 1
    put(ws, row, 2, 'Executed Date', 'label')
    for j in range(n):
        c = put(ws, row, 6 + j, when.date(), 'date')
        c.number_format = 'mm/dd/yyyy'
        c.alignment = openpyxl.styles.Alignment(textRotation=90, horizontal='center')
    ws.row_dimensions[row].height = 55
    row += 1
    put(ws, row, 2, 'Defect ID', 'label')
    row += 1
    put(ws, row, 2, 'Test name (pytest)', 'label')
    for j, case in enumerate(cases):
        c = put(ws, row, 6 + j, case[0], 'res')
        c.alignment = openpyxl.styles.Alignment(textRotation=90, vertical='bottom', horizontal='center')
    ws.row_dimensions[row].height = 12 + 5.5 * max(len(c[0]) for c in cases)
    # summary row 7, from the matrix below
    ws['A7'] = f'=COUNTIF({first}{prow}:{last}{prow},"P")'
    ws['C7'] = f'=COUNTIF({first}{prow}:{last}{prow},"F")'
    ws['F7'] = '=SUM(O7,-A7,-C7)'
    ws['L7'] = f'=COUNTIF({first}{trow}:{last}{trow},"N")'
    ws['M7'] = f'=COUNTIF({first}{trow}:{last}{trow},"A")'
    ws['N7'] = f'=COUNTIF({first}{trow}:{last}{trow},"B")'
    ws['O7'] = f'=COUNTA({first}9:{last}9)'

# --- Statistics
sx = wb['Statistics']
sx['F4'], sx['F6'], sx['B6'] = CREATOR, when.date(), f'{CODE}_UnitTest_v1.0'
sx['B7'] = f'Management service of the logos workspace: {len(FUNCTIONS)} functions. Whole suite run: see Notes in Cover.'
style_row = [copy(sx.cell(row=12, column=c)._style) for c in range(1, 10)]
for row in sx.iter_rows(min_row=12, max_row=sx.max_row):
    for c in row:
        c.value, c.hyperlink = None, None
for i, f in enumerate(FUNCTIONS):
    r, s = 12 + i, f"'{f['code'][:31]}'"
    vals = [i + 1, f['code'], f'={s}!A7', f'={s}!C7', f'={s}!F7', f'={s}!L7', f'={s}!M7', f'={s}!N7', f'={s}!O7']
    for col, v in enumerate(vals, 1):
        sx.cell(row=r, column=col, value=v)._style = copy(style_row[col - 1])
    sx.cell(row=r, column=2).hyperlink = f"#{s}!A1"
end = 11 + len(FUNCTIONS)
for L in 'FGH':  # N, A and B columns, as narrow as Passed/Failed
    sx.column_dimensions[L].width = 10
t = end + 1
sx.cell(row=t, column=2, value='Sub total')
for col in range(3, 10):
    L = get_column_letter(col)
    sx.cell(row=t, column=col, value=f'=SUM({L}12:{L}{end})')
for k, (label, f_) in enumerate([('Test coverage', f'=(C{t}+D{t})*100/I{t}'), ('Test successful coverage', f'=C{t}*100/I{t}'),
                                 ('Normal case', f'=F{t}*100/I{t}'), ('Abnormal case', f'=G{t}*100/I{t}'),
                                 ('Boundary case', f'=H{t}*100/I{t}')]):
    sx.cell(row=t + 2 + k, column=2, value=label)
    sx.cell(row=t + 2 + k, column=4, value=f_).number_format = '0.0'
    sx.cell(row=t + 2 + k, column=5, value='%')

# the template pie charts read the Sub total row of the sample (row 45); point them at ours
for ch in sx._charts:
    a = ch.anchor
    shift = max(0, t + 8 - a._from.row)  # keep the charts below the coverage rows
    a._from.row += shift
    if hasattr(a, "to"):
        a.to.row += shift
    for se in ch.series:
        ref = se.val.numRef
        ref.f = re.sub(r'\$(\d+):', f'${t}:', ref.f)
        ref.f = re.sub(r'\$(\d+)$', f'${t}', ref.f)
        ref.numCache = None
wb.remove(proto)
wb.active = wb.sheetnames.index('Cover')
wb.save(out)
print(len(FUNCTIONS), 'functions', sum(len(f['cases']) for f in FUNCTIONS), 'test cases')
